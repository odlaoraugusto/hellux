"""
Service do módulo Relatórios (Sprint 10).

Gera arquivos em memória (BytesIO) para exportação - nunca grava nada
em disco no servidor. Cada método devolve os bytes prontos para serem
enviados como resposta HTTP (StreamingResponse) pelo Router.

Fase 1 (fluxo de Exame unificado): os relatórios de "Solicitações" e
"Resultados Parciais" agora leem de `Exame` em vez de
`Solicitacao`/`Cultura`.

Fase 1.6 (white-label): logo e título dos relatórios passam a refletir
o `Tenant` do usuário autenticado (`nome_fantasia`, `subtitulo_cabecalho`,
`logo_url`) em vez da marca "Hellux" fixa - ver `_obter_logo_bytes` e o
fallback genérico usado quando `tenant is None` (caso SUPER_ADMIN).
"""
import io
import logging
import os
import uuid
from datetime import date, datetime

import httpx
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    Image,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlalchemy.orm import Session

from app.models.tenant import Tenant
from app.repositories.paciente_repository import PacienteRepository
from app.services.ccih_service import CCIHService
from app.services.exame_service import ExameService

logger = logging.getLogger(__name__)

CABECALHO_FILL = PatternFill(start_color="0F4C81", end_color="0F4C81", fill_type="solid")
CABECALHO_FONT = Font(color="FFFFFF", bold=True)
TITULO_INSTITUCIONAL_FONT = Font(bold=True, size=14, color="0F4C81")
SUBTITULO_INSTITUCIONAL_FONT = Font(italic=True, size=10, color="6B7280")

LOGO_PATH = os.path.join(os.path.dirname(__file__), "..", "assets", "logo.png")
LOGO_DOWNLOAD_TIMEOUT_SEGUNDOS = 5
NOME_INSTITUICAO_PADRAO = "Hellux"

# Tamanho do logo nos relatórios (mesma dimensão usada no PDF, convertida
# aproximadamente para pixels no Excel - openpyxl trabalha em pixels, não cm).
LOGO_LARGURA_PX = 60
LOGO_ALTURA_PX = 60


def _logo_padrao_bytes() -> bytes | None:
    """Bytes do logo padrão do Hellux (`app/assets/logo.png`), se existir."""
    if not os.path.exists(LOGO_PATH):
        return None
    with open(LOGO_PATH, "rb") as arquivo:
        return arquivo.read()


def _obter_logo_bytes(tenant: Tenant | None) -> bytes | None:
    """
    Resolve os bytes do logo a ser exibido no relatório: baixa
    `tenant.logo_url` se o tenant tiver um configurado, com fallback
    gracioso pro logo padrão do Hellux em qualquer cenário adverso.

    Nunca lança exceção - timeout, URL inválida, resposta de erro HTTP ou
    conteúdo que não é uma imagem decodificável (ex.: a URL devolve uma
    página de erro HTML com status 200) tudo cai silenciosamente no logo
    padrão, exatamente como o `os.path.exists` de antes desta fase: um
    problema no branding jamais pode derrubar a geração do relatório.
    """
    if tenant and tenant.logo_url:
        try:
            resposta = httpx.get(tenant.logo_url, timeout=LOGO_DOWNLOAD_TIMEOUT_SEGUNDOS)
            resposta.raise_for_status()
            conteudo = resposta.content
            # Valida que o que veio da URL é mesmo uma imagem antes de
            # deixar chegar no reportlab/openpyxl - `verify()` lança se
            # não conseguir decodificar.
            PILImage.open(io.BytesIO(conteudo)).verify()
            return conteudo
        except Exception:
            logger.warning(
                "Falha ao baixar logo do tenant %s (logo_url=%r) - usando logo padrão.",
                tenant.id,
                tenant.logo_url,
                exc_info=True,
            )

    return _logo_padrao_bytes()


def _estilizar_cabecalho(ws, colunas: list[str]) -> None:
    ws.append(colunas)
    linha_cabecalho = ws.max_row
    for cell in ws[linha_cabecalho]:
        cell.fill = CABECALHO_FILL
        cell.font = CABECALHO_FONT
    for i, coluna in enumerate(colunas, start=1):
        ws.column_dimensions[chr(64 + i)].width = max(18, len(coluna) + 4)


def _adicionar_cabecalho_institucional(ws, tenant: Tenant | None, num_colunas: int) -> None:
    """
    Insere, antes do cabeçalho de colunas, o nome do tenant (+ subtítulo,
    se preenchido) e o logo institucional - mesmo branding do PDF do
    CCIH, adaptado pro Excel (que não tinha nenhum até a Fase 1.6).

    O logo fica ancorado em "A1" (célula da primeira coluna/primeira
    linha); o texto do título começa na coluna seguinte para não ficar
    sobreposto ao logo.
    """
    nome_instituicao = tenant.nome_fantasia if tenant else NOME_INSTITUICAO_PADRAO

    ws.append([None, nome_instituicao])
    linha_titulo = ws.max_row
    ws.cell(row=linha_titulo, column=2).font = TITULO_INSTITUCIONAL_FONT
    if num_colunas > 1:
        ws.merge_cells(start_row=linha_titulo, start_column=2, end_row=linha_titulo, end_column=num_colunas)

    if tenant and tenant.subtitulo_cabecalho:
        ws.append([None, tenant.subtitulo_cabecalho])
        linha_subtitulo = ws.max_row
        ws.cell(row=linha_subtitulo, column=2).font = SUBTITULO_INSTITUCIONAL_FONT
        if num_colunas > 1:
            ws.merge_cells(
                start_row=linha_subtitulo, start_column=2, end_row=linha_subtitulo, end_column=num_colunas
            )

    ws.append([])  # linha em branco separando o cabeçalho institucional do de colunas

    ws.row_dimensions[1].height = 32
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 10)

    logo_bytes = _obter_logo_bytes(tenant)
    if logo_bytes:
        try:
            imagem = XLImage(io.BytesIO(logo_bytes))
            imagem.width = LOGO_LARGURA_PX
            imagem.height = LOGO_ALTURA_PX
            ws.add_image(imagem, "A1")
        except Exception:
            # Mesmo espírito do helper de bytes: um logo que não dá pra
            # anexar (ex.: formato que o openpyxl não reconhece) não pode
            # impedir o resto do relatório de ser gerado.
            logger.warning("Falha ao anexar logo no Excel - relatório seguirá sem logo.", exc_info=True)


class RelatorioService:
    def __init__(self, db: Session):
        self.db = db
        self.paciente_repository = PacienteRepository(db)
        self.ccih_service = CCIHService(db)
        self.exame_service = ExameService(db)

    def gerar_excel_pacientes(self, tenant: Tenant | None = None) -> bytes:
        colunas = ["Prontuário", "Nome", "Setor", "Leito", "Cadastrado em"]
        wb = Workbook()
        ws = wb.active
        ws.title = "Pacientes"
        _adicionar_cabecalho_institucional(ws, tenant, len(colunas))
        _estilizar_cabecalho(ws, colunas)

        pacientes, _ = self.paciente_repository.search(termo=None, skip=0, limit=10_000)
        for p in pacientes:
            ws.append(
                [
                    p.prontuario,
                    p.nome,
                    p.setor or "",
                    p.leito or "",
                    p.created_at.strftime("%d/%m/%Y %H:%M"),
                ]
            )

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    def gerar_excel_exames(self, tenant: Tenant | None = None) -> bytes:
        colunas = [
            "Paciente",
            "Prontuário",
            "Material",
            "Setor",
            "Tipo de Cultura",
            "Status",
            "Data da Coleta",
        ]
        wb = Workbook()
        ws = wb.active
        ws.title = "Exames"
        _adicionar_cabecalho_institucional(ws, tenant, len(colunas))
        _estilizar_cabecalho(ws, colunas)

        exames, _ = self.exame_service.listar(status=None, page=1, page_size=10_000)
        for e in exames:
            ws.append(
                [
                    e.paciente.nome if e.paciente else "",
                    e.paciente.prontuario if e.paciente else "",
                    e.material.nome if e.material else "",
                    e.setor.nome if e.setor else "",
                    e.tipo_cultura.nome if e.tipo_cultura else "",
                    e.status.value,
                    e.data_coleta.strftime("%d/%m/%Y"),
                ]
            )

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    def gerar_excel_exames_parciais(self, tenant: Tenant | None = None) -> bytes:
        colunas = [
            "Paciente",
            "Prontuário",
            "Material",
            "Tipo de Cultura",
            "Status Atual",
            "Microrganismo(s)",
            "Previsão de Liberação",
            "Pendência",
        ]
        wb = Workbook()
        ws = wb.active
        ws.title = "Resultados Parciais"
        _adicionar_cabecalho_institucional(ws, tenant, len(colunas))
        _estilizar_cabecalho(ws, colunas)

        itens = self.exame_service.resultados_parciais()
        for exame, pendencia in itens:
            paciente = exame.paciente
            nomes_micro = (
                ", ".join(i.microrganismo.nome for i in exame.isolados)
                if exame.isolados
                else "—"
            )
            previsao = (
                exame.previsao_liberacao.strftime("%d/%m/%Y")
                if exame.previsao_liberacao
                else "—"
            )
            ws.append(
                [
                    paciente.nome if paciente else "",
                    paciente.prontuario if paciente else "",
                    exame.material.nome if exame.material else "",
                    exame.tipo_cultura.nome if exame.tipo_cultura else "",
                    exame.status.value,
                    nomes_micro,
                    previsao,
                    pendencia,
                ]
            )

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    def gerar_pdf_ccih(
        self,
        data_inicio: date | None,
        data_fim: date | None,
        tenant: Tenant | None = None,
        setor_id: uuid.UUID | None = None,
        vigilancia: bool = False,
    ) -> bytes:
        indicadores = (
            self.ccih_service.indicadores_vigilancia(data_inicio, data_fim, setor_id=setor_id)
            if vigilancia
            else self.ccih_service.indicadores(data_inicio, data_fim, setor_id=setor_id)
        )
        nome_instituicao = tenant.nome_fantasia if tenant else NOME_INSTITUICAO_PADRAO
        titulo = (
            f"{nome_instituicao} — Relatório CCIH (Cultura de Vigilância)"
            if vigilancia
            else f"{nome_instituicao} — Relatório CCIH"
        )
        styles = getSampleStyleSheet()

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            leftMargin=2 * cm,
            rightMargin=2 * cm,
            topMargin=2 * cm,
            bottomMargin=2 * cm,
        )

        cabecalho_texto = [Paragraph(titulo, styles["Title"])]
        if tenant and tenant.subtitulo_cabecalho:
            cabecalho_texto.append(Paragraph(tenant.subtitulo_cabecalho, styles["Normal"]))
        cabecalho_texto += [
            Paragraph(
                f"Período: {indicadores.periodo_inicio.strftime('%d/%m/%Y')} a "
                f"{indicadores.periodo_fim.strftime('%d/%m/%Y')}",
                styles["Normal"],
            ),
            Paragraph(
                f"Setor: {indicadores.filtro_setor or 'Todos os setores'}",
                styles["Normal"],
            ),
            Paragraph(
                f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}",
                styles["Normal"],
            ),
        ]

        logo_bytes = _obter_logo_bytes(tenant)
        if logo_bytes:
            cabecalho = Table(
                [[Image(io.BytesIO(logo_bytes), width=1.6 * cm, height=1.6 * cm), cabecalho_texto]],
                colWidths=[2.2 * cm, None],
            )
            cabecalho.setStyle(
                TableStyle(
                    [
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("LEFTPADDING", (0, 0), (0, 0), 0),
                        ("LEFTPADDING", (1, 0), (1, 0), 10),
                    ]
                )
            )
            elementos = [cabecalho, Spacer(1, 0.8 * cm)]
        else:
            elementos = [*cabecalho_texto, Spacer(1, 0.8 * cm)]

        resumo_dados = [
            ["Solicitações no período", str(indicadores.total_solicitacoes)],
            ["Culturas positivas", str(indicadores.total_culturas_positivas)],
            ["Taxa de positividade", f"{indicadores.taxa_positividade}%"],
        ]
        elementos.append(Paragraph("Resumo Geral", styles["Heading2"]))
        elementos.append(_tabela(resumo_dados, ["Indicador", "Valor"]))
        elementos.append(Spacer(1, 0.6 * cm))

        elementos.append(Paragraph("Distribuição por Setor", styles["Heading2"]))
        if indicadores.distribuicao_por_setor:
            dados_setor = [[d.setor, str(d.total_positivas)] for d in indicadores.distribuicao_por_setor]
            elementos.append(_tabela(dados_setor, ["Setor", "Culturas Positivas"]))
        else:
            elementos.append(Paragraph("Sem dados no período.", styles["Normal"]))
        elementos.append(Spacer(1, 0.6 * cm))

        elementos.append(Paragraph("Perfil Microbiológico", styles["Heading2"]))
        if indicadores.perfil_microbiologico:
            dados_perfil = [
                [p.microrganismo, str(p.quantidade), f"{p.percentual}%"]
                for p in indicadores.perfil_microbiologico
            ]
            elementos.append(_tabela(dados_perfil, ["Microrganismo", "Quantidade", "% do total"]))
        else:
            elementos.append(Paragraph("Sem dados no período.", styles["Normal"]))
        elementos.append(Spacer(1, 0.6 * cm))

        elementos.append(Paragraph("Mapa de Resistência", styles["Heading2"]))
        if indicadores.taxa_resistencia:
            dados_resist = [
                [
                    r.antimicrobiano,
                    str(r.total_testado),
                    str(r.total_resistente),
                    f"{r.percentual_resistente}%",
                    str(r.total_sensivel),
                    f"{r.percentual_sensivel}%",
                ]
                for r in indicadores.taxa_resistencia
            ]
            elementos.append(
                _tabela(
                    dados_resist,
                    [
                        "Antimicrobiano",
                        "Testados",
                        "Resistentes",
                        "% Resistência",
                        "Sensíveis",
                        "% Sensibilidade",
                    ],
                )
            )
        else:
            elementos.append(Paragraph("Nenhum antibiograma liberado no período.", styles["Normal"]))

        doc.build(elementos)
        return buffer.getvalue()


def _tabela(dados: list[list[str]], cabecalho: list[str]) -> Table:
    linhas = [cabecalho] + dados
    tabela = Table(linhas, hAlign="LEFT")
    tabela.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F4C81")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    return tabela
