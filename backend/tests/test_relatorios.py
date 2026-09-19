"""
Testes do módulo Relatórios (Sprint 10).

Como os endpoints retornam arquivos binários (não o contrato JSON
padrão), os testes verificam: status HTTP, content-type, headers de
download e a assinatura binária do arquivo gerado (magic bytes).

Fase 1 (fluxo de Exame unificado): `/api/relatorios/solicitacoes.xlsx`
vira `/api/relatorios/exames.xlsx` e `/api/relatorios/culturas-
parciais.xlsx` vira `/api/relatorios/exames-parciais.xlsx`.

Fase 1.6 (white-label): os testes abaixo cobrem o branding por tenant -
nome/subtítulo aparecendo nos relatórios e o fallback gracioso quando o
`logo_url` do tenant não é acessível. O PDF do CCIH é gerado com o
conteúdo do stream comprimido pelo reportlab (não há utilitário de
extração de texto de PDF neste projeto), então a verificação de que o
tenant certo foi usado é feita "espionando" a chamada ao Service (em vez
de abrir o PDF e procurar o texto); no Excel, como o conteúdo das
células é lido facilmente via `openpyxl`, o teste confirma o nome do
tenant direto na planilha gerada.
"""
import io

import httpx
import openpyxl

from app.services.relatorio_service import RelatorioService
from tests.helpers import criar_exame, criar_microrganismo


def test_exportar_pacientes_excel(authenticated_client):
    authenticated_client.post("/api/pacientes", json={"nome": "Paciente Relatório", "prontuario": "r1"})

    response = authenticated_client.get("/api/relatorios/pacientes.xlsx")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in response.headers["content-disposition"]
    # arquivos .xlsx são arquivos ZIP - começam com a assinatura "PK"
    assert response.content[:2] == b"PK"


def test_exportar_exames_excel(authenticated_client):
    criar_exame(authenticated_client, prontuario="r2")

    response = authenticated_client.get("/api/relatorios/exames.xlsx")

    assert response.status_code == 200
    assert response.content[:2] == b"PK"


def test_exportar_exames_parciais_excel(authenticated_client):
    criar_exame(authenticated_client, prontuario="rp1")  # status padrão: AGUARDANDO_TRIAGEM

    response = authenticated_client.get("/api/relatorios/exames-parciais.xlsx")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.content[:2] == b"PK"


def test_exportar_ccih_pdf(authenticated_client):
    microrganismo = criar_microrganismo(authenticated_client, nome="Micro Relatório")
    criar_exame(
        authenticated_client,
        prontuario="r3",
        status="POSITIVO",
        isolados=[{"microrganismo_id": microrganismo["id"]}],
    )

    response = authenticated_client.get("/api/relatorios/ccih.pdf")

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert "attachment" in response.headers["content-disposition"]
    # arquivos PDF começam com a assinatura "%PDF"
    assert response.content[:4] == b"%PDF"


def test_exportar_ccih_pdf_com_periodo_customizado(authenticated_client):
    response = authenticated_client.get(
        "/api/relatorios/ccih.pdf",
        params={"data_inicio": "2026-01-01", "data_fim": "2026-01-31"},
    )
    assert response.status_code == 200
    assert response.content[:4] == b"%PDF"


def test_excel_pacientes_traz_nome_e_subtitulo_do_tenant(authenticated_client, tenant, db_session):
    tenant.subtitulo_cabecalho = "Rede Hospitalar Teste"
    db_session.commit()

    response = authenticated_client.get("/api/relatorios/pacientes.xlsx")

    assert response.status_code == 200
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    ws = workbook.active
    valores_cabecalho = [
        cell.value for linha in ws.iter_rows(min_row=1, max_row=3) for cell in linha
    ]
    assert tenant.nome_fantasia in valores_cabecalho
    assert tenant.subtitulo_cabecalho in valores_cabecalho


def test_logo_url_invalida_nao_quebra_relatorio_excel(
    authenticated_client, tenant, db_session, monkeypatch
):
    """
    `logo_url` apontando para algo inacessível (rede fora do ar, URL
    inválida, timeout etc.) nunca pode derrubar a geração do relatório -
    o helper `_obter_logo_bytes` degrada pro logo padrão do Hellux
    silenciosamente. Simula a falha de rede mockando `httpx.get`.
    """
    tenant.logo_url = "http://logo-inexistente.invalido/logo.png"
    db_session.commit()

    def _get_com_falha(*args, **kwargs):
        raise httpx.ConnectError("simulando falha de rede no teste")

    monkeypatch.setattr("app.services.relatorio_service.httpx.get", _get_com_falha)

    response = authenticated_client.get("/api/relatorios/pacientes.xlsx")

    assert response.status_code == 200
    assert response.content[:2] == b"PK"
    # o arquivo continua válido/abrível mesmo com a falha no download do logo
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    assert workbook.active is not None


def test_logo_url_invalida_nao_quebra_relatorio_pdf(
    authenticated_client, tenant, db_session, monkeypatch
):
    tenant.logo_url = "http://logo-inexistente.invalido/logo.png"
    db_session.commit()

    def _get_com_falha(*args, **kwargs):
        raise httpx.ConnectError("simulando falha de rede no teste")

    monkeypatch.setattr("app.services.relatorio_service.httpx.get", _get_com_falha)

    response = authenticated_client.get("/api/relatorios/ccih.pdf")

    assert response.status_code == 200
    assert response.content[:4] == b"%PDF"


def test_pdf_ccih_e_gerado_com_o_tenant_do_usuario_autenticado(authenticated_client, tenant, monkeypatch):
    """
    Não existe, neste projeto, um utilitário de extração de texto de PDF
    em teste (o conteúdo do stream sai comprimido pelo reportlab) - a
    verificação de que o título/nome do tenant é usado na geração é
    feita confirmando que o Service recebeu o tenant certo, além do
    200 + content-type esperados (ver docstring do módulo).
    """
    chamadas = {}
    original = RelatorioService.gerar_pdf_ccih

    def _espiao(self, *args, **kwargs):
        chamadas["tenant"] = kwargs.get("tenant")
        return original(self, *args, **kwargs)

    monkeypatch.setattr(RelatorioService, "gerar_pdf_ccih", _espiao)

    response = authenticated_client.get("/api/relatorios/ccih.pdf")

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert chamadas["tenant"] is not None
    assert chamadas["tenant"].id == tenant.id
    assert chamadas["tenant"].nome_fantasia == tenant.nome_fantasia


def test_relatorio_com_super_admin_usa_fallback_generico_sem_tenant(super_admin_client):
    """
    SUPER_ADMIN não pertence a nenhum tenant (`usuario.tenant_id is None`)
    - `get_current_tenant` devolve `None` nesse caso em vez de levantar
    erro, e o relatório cai no fallback genérico (nome "Hellux", sem
    logo do tenant) em vez de quebrar a geração.
    """
    response = super_admin_client.get("/api/relatorios/ccih.pdf")

    assert response.status_code == 200
    assert response.content[:4] == b"%PDF"


def test_downloads_nao_usam_nome_hellux_no_arquivo(authenticated_client):
    """
    Nomes de arquivo de download não devem carregar a marca do
    fornecedor do software (espírito white-label) - ver Fase 1.6.
    """
    respostas = {
        "/api/relatorios/pacientes.xlsx": authenticated_client.get("/api/relatorios/pacientes.xlsx"),
        "/api/relatorios/exames.xlsx": authenticated_client.get("/api/relatorios/exames.xlsx"),
        "/api/relatorios/ccih.pdf": authenticated_client.get("/api/relatorios/ccih.pdf"),
    }
    for resposta in respostas.values():
        assert "hellux" not in resposta.headers["content-disposition"].lower()
